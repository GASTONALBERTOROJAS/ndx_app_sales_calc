"""
Sales Calculation - Tower Component Cost Extractor

Extrae datos de costos de componentes de torres eólicas desde un archivo Excel.

Combina información de dos hojas:
  1. TS SC v26.1: Estructura matricial con datos de torres de acero (400 torres)
  2. TCS_MB SC v26.1: Tabla plana de torres de concreto/híbridas (108 torres)

Produce un archivo Excel unificado con columnas:
  - Component_Category (fijo: "Tower")
  - Component (nombre del componente)
  - Key (identificador de la torre)
  - Brand (marca: Nx para Nordex)
  - Year_Production (año: 2026, 2027, 2028)
  - Region (región canonicalizada: 11 opciones)
  - Cost_EUR (costo en Euros)
  - Cost_USD (costo en Dólares Americanos)

Configuración:
  - Lee config.json del directorio raíz del proyecto
  - Define rutas, años, regiones y nombres de hojas

Uso:
  python scripts/extract_tower_data.py
  
Salida:
  output/SalesCalc_DDMMYY.xlsx (133,588 filas)
  Ejemplo: SalesCalc_220426.xlsx (22 de abril de 2026)
"""

import re
import json
import logging
import warnings
import openpyxl
import pandas as pd
from pathlib import Path
from datetime import datetime

warnings.filterwarnings("ignore", category=FutureWarning, message=".*ChainedAssignmentError.*")

# --- Configuration -----------------------------------------------------------

# Load configuration from config.json
PROJECT_ROOT = Path(__file__).parent.parent
CONFIG_FILE = PROJECT_ROOT / "config.json"

with open(CONFIG_FILE, "r") as f:
    config = json.load(f)

INPUT_FILE = config["input_file_path"]

# Generate dynamic output filename: SalesCalc_DDMMYY.xlsx
today = datetime.now()
date_str = today.strftime("%d%m%y")  # Format: DDMMYY
output_filename = f"SalesCalc_{date_str}.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "output" / output_filename
SHEET_NAME = config["source_sheets"]["ts_sc"]
SHEET_TCS_MB = config["source_sheets"]["tcs_mb"]

TARGET_YEARS = [str(y) for y in config["target_years"]]
CANONICAL_REGIONS = config["canonical_regions"]

# Map raw region strings (from row 7 / row 3) to canonical region
REGION_MAP = {
    "Europe": "Europe",
    "Germany": "Germany",
    "Turkey": "Turkey",
    "Turkey_DOM": "Turkey_DOM",
    "Asia": "Asia",
    "Asia EUR": "Asia",
    "Asia USD": "Asia",
    "China": "China",
    "China EUR": "China",
    "China USD": "China",
    "Poland": "Poland",
    "Italy": "Italy",
    "Greece": "Greece",
    "CAN": "CAN",
    "US": "US",
    "Arcosa US": "US",
    "Arcosa(US)": "US",
    "CS Wind US": "US",
    "CS Wind(US)": "US",
}

# Components that appear WITH year+region in row 3
COMPONENTS_WITH_YEAR = [
    "tower shell",
    "Tower internals",
    "Anchor cage",
    "Tower bolts set",
]

# Components WITHOUT year but WITH region in row 3
COMPONENTS_REGION_ONLY = [
    "Steel Tower Quality Inspectors",
]

# Components WITHOUT year and WITHOUT region (single column each)
COMPONENTS_GLOBAL = [
    "Option Coating c4/c5",
    "Option Anchor Cage: German requirements for NAT AC",
    "Option hybrid tower: MB Monthly Cost Indexation",
    "Option hybrid tower: white coating of concrete part",
    "Option hybrid tower: no red stripe on concrete part",
]

# --- Logging -----------------------------------------------------------------

logging.basicConfig(
    level=logging.WARNING,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
log = logging.getLogger(__name__)

# --- Helpers -----------------------------------------------------------------

def parse_row3_value(cell_value: str):
    """
    Parse a concatenated row-3 string like:
      'tower shell2026Cost_Currency1Europe'
    Returns (component, year, currency_type, region) or None.

    currency_type is 1 or 2 (from Cost_Currency1 / Cost_Currency2).
    year and region may be None for global/region-only components.
    """
    if not cell_value or not isinstance(cell_value, str):
        return None

    # Try to extract currency type
    m_curr = re.search(r"Cost_Currency([12])", cell_value)
    if not m_curr:
        return None
    currency_type = int(m_curr.group(1))

    # Part before Cost_Currency contains component + optional year
    before_cost = cell_value[: m_curr.start()]
    # Part after Cost_CurrencyX is the region
    after_cost = cell_value[m_curr.end():]
    region = after_cost.strip() if after_cost.strip() else None

    # Try to extract year from the end of before_cost
    m_year = re.search(r"(2025|2026|2027|2028)$", before_cost)
    if m_year:
        year = m_year.group(1)
        component = before_cost[: m_year.start()]
    else:
        year = None
        component = before_cost

    return component.strip(), year, currency_type, region


def extract_region_from_row7(cell_value: str):
    """
    Row 7 values like 'tower shell\\nEurope' — extract the region part.
    """
    if not cell_value or not isinstance(cell_value, str):
        return None
    parts = cell_value.split("\n")
    if len(parts) >= 2:
        return parts[-1].strip()
    return None


def normalize_region(raw_region: str) -> str | None:
    """
    Map a raw region string to one of the 11 canonical regions.
    Returns None if the region is not recognized.
    """
    if not raw_region:
        return None
    mapped = REGION_MAP.get(raw_region)
    if mapped is None:
        log.warning("Unknown region '%s' — skipping.", raw_region)
    return mapped


# --- Reporting ---------------------------------------------------------------

def _print_summary(df: pd.DataFrame, output_path: Path) -> None:
    """Print extraction summary: years, regions, components, data quality."""
    sep = "=" * 60
    total = len(df)
    ncols = len(df.columns)

    years = sorted(df["Year_Production"].dropna().unique())
    regions = sorted(df["Region"].dropna().unique())

    comp_stats = (
        df.groupby("Component")
        .agg(
            Filas=("Component", "count"),
            Con_EUR=("Cost_EUR", lambda x: x.notna().sum()),
            Con_USD=("Cost_USD", lambda x: x.notna().sum()),
        )
        .reset_index()
        .sort_values("Filas", ascending=False)
    )

    con_eur = df["Cost_EUR"].notna().sum()
    con_usd = df["Cost_USD"].notna().sum()
    con_ambas = (df["Cost_EUR"].notna() & df["Cost_USD"].notna()).sum()

    print(f"\n{sep}")
    print("EXTRACCION COMPLETADA")
    print(sep)

    print(f"\n  Archivo guardado:")
    print(f"    {output_path}")
    print(f"    {total:,} filas | {ncols} columnas")

    print(f"\n  Years:")
    for y in years:
        print(f"    - {int(y)}")

    print(f"\n  Regiones ({len(regions)}):")
    for r in regions:
        print(f"    - {r}")

    print(f"\n  Componentes encontrados ({len(comp_stats)} unicos):\n")
    hdr = f"    {'Componente':<50} {'Filas':>7}"
    print(hdr)
    print("    " + "-" * (len(hdr) - 4))
    for _, row in comp_stats.iterrows():
        print(f"    {row['Component']:<50} {int(row['Filas']):>7,}")

    print(f"\n  Calidad de datos:")
    print(f"    Filas con Cost_EUR:    {con_eur:>7,}  ({con_eur/total*100:.1f}%)")
    print(f"    Filas con Cost_USD:    {con_usd:>7,}  ({con_usd/total*100:.1f}%)")
    print(f"    Filas con ambas:       {con_ambas:>7,}")

    print(f"\n{sep}\n")

# --- Main extraction ---------------------------------------------------------

def main():
    log.info("Opening workbook: %s", INPUT_FILE)
    wb = openpyxl.load_workbook(INPUT_FILE, data_only=True, read_only=True)
    ws = wb[SHEET_NAME]
    log.info("Sheet '%s' loaded successfully.", SHEET_NAME)

    # ---- Step 1: Build column map from rows 1-3 and 7 ----------------------
    # For each target column, we store:
    #   col_index -> {component, year, currency_type, region}

    # Read rows 1, 2, 3, 7 into dicts keyed by column index
    row_data = {1: {}, 2: {}, 3: {}, 7: {}}
    for row in ws.iter_rows(min_row=1, max_row=7):
        for c in row:
            try:
                col = c.column
                r = c.row
                if r in row_data and c.value is not None:
                    row_data[r][col] = c.value
            except AttributeError:
                # EmptyCell in read_only mode
                continue

    log.info(
        "Header rows loaded — row1: %d, row2: %d, row3: %d, row7: %d cells",
        len(row_data[1]), len(row_data[2]), len(row_data[3]), len(row_data[7]),
    )

    # Parse row 3 to identify target columns
    all_target_components = (
        [c.lower() for c in COMPONENTS_WITH_YEAR]
        + [c.lower() for c in COMPONENTS_REGION_ONLY]
        + [c.lower() for c in COMPONENTS_GLOBAL]
    )

    col_map = {}  # col_index -> dict with keys: component, year, currency_type, region
    for col, val in row_data[3].items():
        parsed = parse_row3_value(str(val))
        if parsed is None:
            continue
        component, year, currency_type, region = parsed

        # Check if this component is one we care about
        if component.lower() not in all_target_components:
            continue

        # Filter: for COMPONENTS_WITH_YEAR, only keep target years
        if component.lower() in [c.lower() for c in COMPONENTS_WITH_YEAR]:
            if year not in TARGET_YEARS:
                continue

        # For region, prefer row 7 (cleaner) over row 3
        region_r7 = extract_region_from_row7(row_data[7].get(col))
        raw_region = region_r7 or region
        final_region = normalize_region(raw_region)

        # Skip columns with unrecognized regions
        if raw_region and final_region is None:
            continue

        col_map[col] = {
            "component": component,
            "year": year,
            "currency_type": currency_type,
            "region": final_region,
        }

    log.info("Identified %d target columns to extract.", len(col_map))

    # ---- Step 2: Read data rows (row 8+) ------------------------------------
    records = []
    row_count = 0

    for row in ws.iter_rows(min_row=8, max_row=1000):
        # Build a dict of col -> value for this row
        row_vals = {}
        for c in row:
            try:
                row_vals[c.column] = c.value
            except AttributeError:
                continue

        key = row_vals.get(4)  # Column D = Key (tower name)
        if not key:
            continue

        brand = row_vals.get(5)  # Column E = Brand
        row_count += 1

        # For each target column, extract the value
        for col, meta in col_map.items():
            value = row_vals.get(col)

            # Components without year in source: replicate for each target year
            if meta["year"] is None:
                years_to_emit = TARGET_YEARS
            else:
                years_to_emit = [meta["year"]]

            # Global components (no region): replicate for each canonical region
            if meta["region"] is None:
                regions_to_emit = CANONICAL_REGIONS
            else:
                regions_to_emit = [meta["region"]]

            for yr in years_to_emit:
                for rgn in regions_to_emit:
                    records.append({
                        "Component_Category": "Tower",
                        "Component": meta["component"],
                        "Key": key,
                        "Brand": brand,
                        "Year_Production": yr,
                        "Region": rgn,
                        "currency_type": meta["currency_type"],
                        "value": value,
                    })

    log.info("Extracted %d raw records from %d tower rows (TS SC).", len(records), row_count)
    wb.close()

    # ---- Step 2b: Extract from TCS_MB SC v26.1 (flat table) ----------------
    log.info("Opening workbook again for TCS_MB extraction...")
    wb2 = openpyxl.load_workbook(INPUT_FILE, data_only=True, read_only=True)
    ws2 = wb2[SHEET_TCS_MB]

    # Component columns to unpivot: (col_index, component_name)
    TCS_MB_COMPONENTS = [
        (5, "Tower Shell"),
        (6, "Tower Internals"),
        (7, "Foundations"),
        (8, "Concrete Tower Keystones + Internals"),
        (9, "Concrete Tower Logistics"),
        (10, "Concrete Tower C&I"),
        (11, "Tower Bolts Set"),
        (13, "Option Coating c4/c5"),
        (14, "Steel Tower Quality Inspectors"),
        (15, "Option hybrid tower: white coating of concrete part"),
        (16, "Option hybrid tower: no red stripe on concrete part"),
    ]

    tcs_records = []
    tcs_row_count = 0

    for row in ws2.iter_rows(min_row=86, max_row=500):
        row_vals = {}
        for c in row:
            try:
                row_vals[c.column] = c.value
            except AttributeError:
                continue

        key = row_vals.get(2)   # Col B = KEY
        if not key:
            continue

        sourcing = row_vals.get(3)  # Col C = Sourcing (Region)
        year = row_vals.get(4)      # Col D = Year
        tcs_row_count += 1

        for col_idx, comp_name in TCS_MB_COMPONENTS:
            value = row_vals.get(col_idx)
            tcs_records.append({
                "Component_Category": "Tower",
                "Component": comp_name,
                "Key": key,
                "Brand": "Nx",
                "Year_Production": year,
                "Region": sourcing,
                "currency_type": 1,
                "value": value,
            })

    log.info(
        "Extracted %d raw records from %d tower rows (TCS_MB).",
        len(tcs_records), tcs_row_count,
    )
    wb2.close()

    # Combine both sources
    records.extend(tcs_records)
    log.info("Combined total: %d raw records.", len(records))

    # ---- Step 3: Pivot Currency1 / Currency2 into separate columns ----------
    df = pd.DataFrame(records)

    # Normalize component names to consistent Title Case
    COMPONENT_NAME_MAP = {
        "tower shell": "Tower Shell",
        "Tower internals": "Tower Internals",
        "Tower bolts set": "Tower Bolts Set",
    }
    df.loc[:, "Component"] = df["Component"].replace(COMPONENT_NAME_MAP)

    # Normalize Year_Production to int
    df.loc[:, "Year_Production"] = pd.to_numeric(df["Year_Production"], errors="coerce").astype("Int64")

    # Create pivot columns
    df_c1 = df[df["currency_type"] == 1].rename(columns={"value": "Cost_EUR"})
    df_c2 = df[df["currency_type"] == 2].rename(columns={"value": "Cost_USD"})

    merge_keys = [
        "Component_Category", "Component", "Key", "Brand",
        "Year_Production", "Region",
    ]

    # Some regions only have Currency1, others have both
    df_c1 = df_c1.drop(columns=["currency_type"])
    df_c2 = df_c2[merge_keys + ["Cost_USD"]]

    df_final = pd.merge(df_c1, df_c2, on=merge_keys, how="outer")

    # Clean up: sort and reorder columns
    output_cols = [
        "Component_Category", "Component", "Key", "Brand",
        "Year_Production", "Region", "Cost_EUR", "Cost_USD",
    ]
    df_final = df_final[output_cols].sort_values(
        by=["Component", "Year_Production", "Region", "Key"]
    ).reset_index(drop=True)

    log.info("Final dataframe: %d rows, %d columns.", len(df_final), len(df_final.columns))

    # ---- Step 4: Write output -----------------------------------------------
    df_final.to_excel(OUTPUT_FILE, index=False, sheet_name="Tower_Extracted")
    log.info("Output written to: %s", OUTPUT_FILE)

    _print_summary(df_final, OUTPUT_FILE)


if __name__ == "__main__":
    main()
