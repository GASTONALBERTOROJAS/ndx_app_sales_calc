import openpyxl, re, os
import pandas as pd
from sqlalchemy import create_engine
from dotenv import load_dotenv

load_dotenv('.env')

db_user = os.getenv('DB_USER', 'postgres')
db_pass = os.getenv('DB_PASS', '')
db_host = os.getenv('DB_HOST', 'localhost')
db_port = os.getenv('DB_PORT', '5432')
db_name = os.getenv('DB_NAME', 'powerbi_reports')
conn_string = f'postgresql://{db_user}:{db_pass}@{db_host}:{db_port}/{db_name}'

# 1. Normalización idéntica a la que aplicamos
def normalize_tower_name(name):
    if not name or str(name).strip().lower() == 'key':
        return name
    name_str = str(name).strip()
    normed = re.sub(r'\s+', ' ', name_str).upper()
    def add_suffix(m):
        prefix = m.group(1)
        num = m.group(2)
        return f'{prefix}{num}-00'
    normed = re.sub(r'\b(TS)(\d+)\b(?!-)', add_suffix, normed)
    return normed

# 2. Leer Sales Calc desde Postgres
engine = create_engine(conn_string)
df_sales = pd.read_sql('SELECT DISTINCT "Key" FROM "03_entrega"."salescalc_tower_sales"', con=engine)
sales_keys = set(df_sales['Key'].dropna())

# 3. Leer Forecast desde el Excel original
wb = openpyxl.load_workbook(r'C:\Users\RojasG1\OneDrive - Nordex SE\Desktop\book3\Book3_marked_backup.xlsx', data_only=True)
ws = wb['Sheet1']

forecast_towers = set()
for r in range(3, ws.max_row+1):
    v = ws.cell(r, 15).value  # Columna Full Tower info
    if v and str(v).strip() and str(v).strip().lower() != 'full tower info':
        norm_v = normalize_tower_name(str(v))
        forecast_towers.add((str(v).strip(), norm_v))

# 4. Comparar
matches = []
no_matches = []

for original, norm in forecast_towers:
    if norm in sales_keys:
        matches.append((original, norm))
    else:
        no_matches.append((original, norm))

print('=== RESULTADOS DE LA COMPARACIÓN ===')
print(f'Total torres únicas en Forecast: {len(forecast_towers)}')
print(f'Match exitoso con BD (PostgreSQL): {len(matches)}')
print(f'Sin match: {len(no_matches)}')

if no_matches:
    print('\n=== TORRES SIN MATCH ===')
    for orig, norm in sorted(no_matches):
        print(f'  Forecast Original: [{orig}]')
        print(f'  Buscado como:      [{norm}]')
        print('  ---')
